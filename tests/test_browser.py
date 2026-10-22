"""Real Chromium + Pyodide import/export smoke check; requires CDN access."""

import io
import base64
import threading
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
from tests.test_groww import payload, stock_book, stock_row, mf_book, mf_row
from tests.test_export_scale import scale_input

ROOT = Path(__file__).resolve().parents[1]


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *_args):
        pass


@pytest.mark.browser
def test_browser_import_download_and_reset(tmp_path):
    server = ThreadingHTTPServer(("127.0.0.1", 0), partial(QuietHandler, directory=str(ROOT)))
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={"width": 1440, "height": 1000})
            failures = []
            requests = []
            page.on("pageerror", lambda error: failures.append(str(error)))
            page.on("request", lambda request: requests.append((request.url, request.method, request.post_data)))
            page.goto(f"http://127.0.0.1:{server.server_port}", wait_until="networkidle")
            csp = page.locator('meta[http-equiv="Content-Security-Policy"]').get_attribute("content")
            assert csp == (
                "default-src 'self'; script-src 'self' https://cdn.jsdelivr.net 'wasm-unsafe-eval'; "
                "connect-src 'self' https://cdn.jsdelivr.net; style-src 'self'; img-src 'self'; "
                "object-src 'none'; base-uri 'self'; form-action 'self'"
            )
            expect(page.locator("footer")).to_contain_text("v0.3.1")
            cold_file = payload(stock_book({"Short Term trades": [stock_row()]}), "first-native.xlsx")
            page.locator("#files").set_input_files({
                "name": cold_file["name"], "mimeType": "application/octet-stream",
                "buffer": base64.b64decode(cold_file["base64"]),
            })
            page.locator("#review").click()
            expect(page.locator("#broker-stcg")).to_have_text("100.20", timeout=180000)
            pyodide_script = page.locator(
                'script[src="https://cdn.jsdelivr.net/pyodide/v0.26.2/full/pyodide.js"]'
            )
            assert pyodide_script.get_attribute("integrity") == (
                "sha384-tVslJOEkg7nVRW3Y3/ReGX0NnonNrbcmt1R5qFbQXQdGa2chRkoJYHAjAsv3zoTq"
            )
            assert pyodide_script.get_attribute("crossorigin") == "anonymous"
            page.locator("#reset").click()
            page.locator("#files").set_input_files(str(ROOT / "samples" / "canonical-demo.csv"))
            expect(page.locator("#review")).to_be_enabled()
            page.locator("#review").click()
            expect(page.locator("#results")).to_be_visible(timeout=180000)
            expect(page.locator("#stcg")).to_have_text("1965.20")
            expect(page.locator("#ltcg")).to_have_text("7500.00")
            expect(page.locator("#blocked")).to_have_text("1")
            expect(page.locator("#lots tr")).to_have_count(4)
            with page.expect_download(timeout=120000) as download_event:
                page.locator("#export").click()
            download = download_event.value
            assert download.suggested_filename == "ITR-Stock-Assist-FY-2025-26-INCOMPLETE.xlsx"
            content = Path(download.path()).read_bytes()
            book = load_workbook(io.BytesIO(content))
            assert book["Summary"]["B7"].value == "1965.20"
            assert book["Source rows"].max_row == 5
            assert len(book.sheetnames) == 9
            assert book.sheetnames[0] == "Filing review"
            assert book["Filing review"]["B3"].value == "2026-27"
            assert "FULL_FY" in page.locator("#filing-year").inner_text()
            expect(page.locator("#filing-groups")).to_contain_text("READY_WORKING")
            assert not failures
            assert all(method == "GET" and body is None for _, method, body in requests)
            assert all(url.startswith(("http://127.0.0.1:", "https://cdn.jsdelivr.net/")) for url, _, _ in requests)
            page.screenshot(path=str(tmp_path / "desktop.png"), full_page=True)
            page.set_viewport_size({"width": 390, "height": 844})
            assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
            page.screenshot(path=str(tmp_path / "mobile.png"), full_page=True)
            page.locator("#fy").select_option("2024-25")
            expect(page.locator("#results")).to_be_hidden()
            expect(page.locator("#export")).to_be_disabled()
            page.locator("#review").click()
            expect(page.locator("#results")).to_be_visible(timeout=30000)
            expect(page.locator("#included")).to_have_text("0")
            expect(page.locator("#blocked")).to_have_text("4")
            expect(page.locator("#filing-year")).to_contain_text("before 23 July 2024")
            expect(page.locator("#filing-groups tr")).to_have_count(0)
            page.locator("#reset").click()
            expect(page.locator("#results")).to_be_hidden()
            expect(page.locator("#review")).to_be_disabled()
            page.locator("#sample").click()
            expect(page.locator("#review")).to_be_enabled()
            expect(page.locator("#fy")).to_have_value("2025-26")
            page.locator("#files").set_input_files({
                "name": "wrong.csv", "mimeType": "text/csv", "buffer": b"name,value\nx,2\n",
            })
            page.locator("#review").click()
            expect(page.locator("#error")).to_have_text(
                "wrong.csv: headers must match the canonical template exactly. "
                "Download the sample CSV and map one disposed acquisition lot per row.",
                timeout=30000,
            )
            assert "Traceback" not in page.locator("#error").inner_text()
            expect(page.locator("#results")).to_be_hidden()
            stock_file = payload(stock_book({"Short Term trades": [stock_row()]}), "stock.xlsx")
            fund_file = payload(mf_book({"Equity Category": [mf_row()]}), "fund.xlsx")
            page.locator("#files").set_input_files([{
                "name": f["name"],
                "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "buffer": base64.b64decode(f["base64"]),
            } for f in (stock_file, fund_file)])
            page.locator("#review").click()
            expect(page.locator("#results")).to_be_visible(timeout=30000)
            expect(page.locator("#broker-summary")).to_be_visible()
            expect(page.locator("#broker-stcg")).to_have_text("100.20")
            expect(page.locator("#broker-ltcg")).to_have_text("123.46")
            expect(page.locator("#broker-count")).to_have_text("2")
            expect(page.locator("#stcg")).to_have_text("0.00")
            expect(page.locator("#included")).to_have_text("0")
            expect(page.locator("#lots")).to_contain_text("scheme SYNTH123")
            assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
            page.set_viewport_size({"width": 1440, "height": 1000})
            page.screenshot(path=str(tmp_path / "native-desktop.png"), full_page=True)
            with page.expect_download(timeout=30000) as native_download:
                page.locator("#export").click()
            native_book = load_workbook(io.BytesIO(Path(native_download.value.path()).read_bytes()))
            assert native_book["Source rows"].max_row == 3
            assert native_book["Import manifest"].max_row == 3
            expect(page.locator("#filing-groups")).to_contain_text("PROVISIONAL")
            mapping = list(native_book["Filing map"].values)
            headings = mapping[0]
            mapped = [dict(zip(headings, row)) for row in mapping[1:]]
            assert all(row["status"] != "READY_WORKING" for row in mapped)
            assert all(row["value"] is None for row in mapped if row["field_id"] == "cg.transfer_expenses")
            assert any(row["field_id"] == "112a.isin" and row["value"] is None for row in mapped)
            etf_file = payload(stock_book({"Short Term trades": [
                stock_row(**{"ISIN": "INF000A01001"}),
            ]}), "synthetic-etf.xlsx")
            debt_file = payload(mf_book({"Debt (Specified - Other than Equity) Category": [
                mf_row(Category="Debt"),
            ]}), "synthetic-debt.xlsx")
            intraday_file = payload(stock_book({"Intraday trades": [stock_row()]}),
                                    "synthetic-intraday.xlsx")
            page.locator("#files").set_input_files([{
                "name": f["name"], "mimeType": "application/octet-stream",
                "buffer": base64.b64decode(f["base64"]),
            } for f in (etf_file, debt_file, intraday_file)])
            page.locator("#review").click()
            expect(page.locator("#filing-groups")).to_contain_text(
                "Non-equity (Schedule CG A5, slab rate)", timeout=30000)
            expect(page.locator("#filing-groups")).to_contain_text("Not reported")
            expect(page.locator("#intraday-warning")).to_be_visible()
            expect(page.locator("#intraday-warning")).to_have_text(
                "Intraday = speculative business income; ITR-3 likely required.")
            with page.expect_download(timeout=30000) as non_equity_download:
                page.locator("#export").click()
            non_equity_book = load_workbook(io.BytesIO(
                Path(non_equity_download.value.path()).read_bytes()))
            assert non_equity_book["Filing groups"]["A2"].value == \
                "Non-equity (Schedule CG A5, slab rate)"
            assert non_equity_book["Filing groups"]["J2"].value is None
            non_equity_map = list(non_equity_book["Filing map"].values)
            assert {row[4] for row in non_equity_map[1:]} >= {
                "cg.consideration", "cg.acquisition_cost", "cg.transfer_expenses",
                "cg.non_equity_net",
            }
            assert any(row[0].value == "Intraday warning" and row[1].value ==
                       "Intraday = speculative business income; ITR-3 likely required."
                       for row in non_equity_book["Summary"].iter_rows())
            assert any(row[0].value == "Intraday warning" and row[1].value ==
                       "Intraday = speculative business income; ITR-3 likely required."
                       for row in non_equity_book["Filing review"].iter_rows())
            zero_file = payload(stock_book({"Short Term trades": [stock_row(**{
                "Buy value": 0, "Sell value": 0, "Realised P&L": 0,
            })]}), "synthetic-zero.xlsx")
            page.locator("#files").set_input_files({
                "name": zero_file["name"], "mimeType": "application/octet-stream",
                "buffer": base64.b64decode(zero_file["base64"]),
            })
            page.locator("#review").click()
            expect(page.locator("#lots")).to_contain_text(
                "likely rights entitlement or corporate action - confirm", timeout=30000)
            with page.expect_download(timeout=30000) as zero_download:
                page.locator("#export").click()
            zero_book = load_workbook(io.BytesIO(Path(zero_download.value.path()).read_bytes()))
            assert any("likely rights entitlement or corporate action - confirm" in row[3].value
                       for row in zero_book["Review issues"].iter_rows(min_row=2))
            dividend_file = payload(stock_book(
                {"Short Term trades": [stock_row()]}, dividends=1234.56,
            ), "synthetic-dividends.xlsx")
            page.locator("#files").set_input_files({
                "name": dividend_file["name"], "mimeType": "application/octet-stream",
                "buffer": base64.b64decode(dividend_file["base64"]),
            })
            page.locator("#review").click()
            expect(page.locator("#dividend-info")).to_be_visible(timeout=30000)
            expect(page.locator("#dividend-info")).to_contain_text("Groww-reported dividends: INR 1234.56")
            expect(page.locator("#dividend-info")).to_contain_text("Report in Schedule OS")
            with page.expect_download(timeout=30000) as dividend_download:
                page.locator("#export").click()
            dividend_book = load_workbook(io.BytesIO(
                Path(dividend_download.value.path()).read_bytes()))
            assert any(row[0].value == "Groww-reported dividends" and
                       "Report in Schedule OS" in row[1].value
                       for row in dividend_book["Filing review"].iter_rows())
            empty_file = payload(stock_book(), "empty.xlsx")
            page.locator("#files").set_input_files({
                "name": empty_file["name"], "mimeType": "application/octet-stream",
                "buffer": base64.b64decode(empty_file["base64"]),
            })
            page.locator("#review").click()
            expect(page.locator("#review-state")).to_contain_text("EMPTY REPORT", timeout=30000)
            expect(page.locator("#manifest")).to_contain_text("no disposal rows")
            page.locator("#files").set_input_files({
                "name": "renamed.xlsx", "mimeType": "application/octet-stream",
                "buffer": base64.b64decode(payload(stock_book(title="P&L Statement for stocks"))["base64"]),
            })
            page.locator("#review").click()
            expect(page.locator("#error")).to_contain_text("P&L reports are not supported", timeout=30000)
            expect(page.locator("#results")).to_be_hidden()
            page.locator("#reset").click()
            assert page.locator("#manifest").text_content() == ""
            assert page.locator("#broker-stcg").text_content() == ""
            assert page.locator("#filing-groups").text_content() == ""
            assert page.locator("#filing-year").text_content() == ""
            scaled = scale_input(1000)
            page.locator("#files").set_input_files({
                "name": scaled["name"], "mimeType": "text/csv",
                "buffer": scaled["text"].encode("utf-8"),
            })
            page.locator("#review").click()
            expect(page.locator("#included")).to_have_text("1000", timeout=30000)
            with page.expect_download(timeout=120000) as scaled_download:
                page.locator("#export").click()
            scaled_book = load_workbook(io.BytesIO(Path(scaled_download.value.path()).read_bytes()))
            assert scaled_book["Disposed lots"].max_row == 1001
            assert scaled_book["Filing map"].max_row == 6001
            assert scaled_book["Source rows"]["C1001"].value == "scale-00999"
            assert scaled_book["Source rows"]["B1001"].value == 1001
            assert scaled_book["Source rows"]["E1001"].data_type == "s"
            assert scaled_book["Summary"]["B7"].value == "-800.00"
            expect(page.locator("#review")).to_be_enabled()
            expect(page.locator("#export")).to_be_enabled()
            assert not failures
            assert all(method == "GET" and body is None for _, method, body in requests)
            assert all(url.startswith(("http://127.0.0.1:", "https://cdn.jsdelivr.net/")) for url, _, _ in requests)
            browser.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


@pytest.mark.browser
def test_forms_ring_dialog_and_responsive_layout():
    server = ThreadingHTTPServer(("127.0.0.1", 0), partial(QuietHandler, directory=str(ROOT)))
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={"width": 1440, "height": 1000})
            failures = []
            console = []
            page.on("pageerror", lambda error: failures.append(str(error)))
            page.on("console", lambda message: console.append((message.type, message.text)))
            page.goto(f"http://127.0.0.1:{server.server_port}", wait_until="networkidle")
            expect(page.locator('[data-ring="home"] .card3')).to_have_count(6)
            assert page.locator("[style]").count() == page.locator(
                ".card3[style], .ring[style], .orb[style], .tilt[style], .deckcard[style], "
                "#lab-fill[style], #lab-mark[style], [data-ring][style], .stage[style]").count()
            page.locator("#sample").click()
            page.locator("#review").click()
            expect(page.locator("#results")).to_be_visible(timeout=180000)
            cards = page.locator(".forms-ring .card3")
            expect(cards).to_have_count(8)
            expect(page.locator('.forms-ring [data-slot="stcg"]')).to_have_text("\u20b91,965.20")
            expect(page.locator('.forms-ring [data-slot="ltcg"]')).to_have_text("\u20b97,500.00")
            expect(page.locator("#kpi-stcg")).to_have_text("1,965.20")
            stage = page.locator('[data-ring="forms"] .stage')
            stage.hover()
            a2 = page.locator('.forms-ring [data-card="stcg"] .cardbtn')
            dialog = page.locator('[role="dialog"]')
            for _ in range(2):
                a2.click()
                page.wait_for_timeout(1000)
                if dialog.is_visible():
                    break
            expect(dialog).to_be_visible()
            assert dialog.get_attribute("aria-modal") == "true"
            expect(page.locator("#dlg-title")).to_contain_text("1,965.20")
            expect(page.locator("#dlg-close")).to_be_focused()
            expect(page.locator("#paperbox .page").first).to_be_visible()
            expect(page.locator("#paperbox .prow.hl")).to_have_count(1)
            expect(page.locator("#listbox .fl").first).to_contain_text("Full value of consideration")
            page.locator("#seg-fill").click()
            assert all(tag in ("FILL", "CHECK") for tag in page.locator("#listbox .tag").all_inner_texts())
            page.set_viewport_size({"width": 390, "height": 844})
            page.wait_for_timeout(300)
            assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
            page.set_viewport_size({"width": 1440, "height": 1000})
            page.keyboard.press("Escape")
            expect(page.locator("#dlg-overlay")).to_be_hidden()
            expect(a2).to_be_focused()
            page.locator('[data-ring="forms"] .dot').nth(3).click()
            stage.hover()
            page.wait_for_timeout(1100)
            page.locator('.forms-ring [data-card="setoff"] .cardbtn').click()
            expect(dialog).to_be_visible()
            assert set(page.locator("#listbox .tag").all_inner_texts()) <= {"CHECK", "SKIP"}
            page.mouse.click(5, 5)
            expect(page.locator("#dlg-overlay")).to_be_hidden()
            page.locator("#reset").click()
            expect(page.locator("#results")).to_be_hidden()
            expect(page.locator('.forms-ring [data-slot="stcg"]')).to_have_text("")
            for width, height in ((390, 844), (820, 1180), (2560, 1440)):
                page.set_viewport_size({"width": width, "height": height})
                page.wait_for_timeout(200)
                assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth"), width
            assert not failures
            assert not [text for kind, text in console if kind == "error" or "Content Security Policy" in text]
            browser.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join()
