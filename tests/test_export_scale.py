"""Check export work by scan counts rather than machine-dependent time limits."""

import csv
import io
from collections import Counter

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from stock_assist.core import HEADERS, review, workbook_bytes
from tests.test_core import SAMPLE


def scale_input(count):
    template = next(csv.DictReader(io.StringIO(SAMPLE)))
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, HEADERS)
    writer.writeheader()
    for index in range(count):
        writer.writerow({
            **template, "lot_id": f"scale-{index:05d}", "asset_name": f"=literal-{index}",
            "proceeds_inr": "0.20", "cost_inr": "1.00", "transfer_expenses_inr": "0",
        })
    return {"name": "synthetic-scale.csv", "text": stream.getvalue()}


@pytest.mark.parametrize("count", [10, 1000])
def test_export_scans_are_bounded_and_all_row_indices_survive(monkeypatch, count):
    data = review([scale_input(count)], "2025-26")
    scans = Counter()
    original_getitem = Worksheet.__getitem__

    def indexed(self, key):
        scans["row_indexing"] += 1
        return original_getitem(self, key)

    with monkeypatch.context() as patch:
        patch.setattr(Worksheet, "__getitem__", indexed)
        for name in ("max_row", "max_column"):
            original = getattr(Worksheet, name)

            def measured(self, getter=original.fget, metric=name):
                scans[metric] += 1
                return getter(self)

            patch.setattr(Worksheet, name, property(measured))
        content = workbook_bytes(data)

    # Whole-sheet bounds/row indexing may be used in fixed formatting passes,
    # never once per appended row. These budgets are independent of lot count.
    assert scans["row_indexing"] <= 9
    assert scans["max_row"] <= 30
    assert scans["max_column"] <= 40
    book = load_workbook(io.BytesIO(content))
    assert len(book.sheetnames) == 9
    assert book["Disposed lots"].max_row == count + 1
    assert book["Source rows"].max_row == count + 1
    assert book["Filing map"].max_row == count * 6 + 1
    for index in range(count):
        row = index + 2
        for name in ("Disposed lots", "Source rows"):
            sheet = book[name]
            assert sheet.cell(row, 2).value == row
            assert sheet.cell(row, 3).value == f"scale-{index:05d}"
            literal = sheet.cell(row, 5)
            assert literal.value == f"=literal-{index}"
            assert literal.data_type == "s"
            assert literal.alignment.wrap_text
        map_row = index * 6 + 2
        assert book["Filing map"].cell(map_row, 4).value == f"lot-{index + 1:06d}"
        assert book["Filing map"].cell(map_row, 16).value == row
        assert book["Filing map"].cell(map_row + 2, 9).value == "0"
        assert book["Filing map"].cell(map_row + 3, 9).value == "-0.80"
        assert book["Filing map"].cell(map_row + 4, 9).value is None
        assert book["Filing map"].cell(map_row + 4, 11).value == "NOT_COMPUTED"
    assert book["Filing map"].freeze_panes == "A2"
    assert book["Filing map"].auto_filter.ref == f"A1:U{count * 6 + 1}"
    assert book["Review issues"].max_row == 1
