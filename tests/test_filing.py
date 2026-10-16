"""AY-specific destinations checked against official V1.2 schemas; see research doc."""

import io
from decimal import Decimal, localcontext

import pytest
from openpyxl import load_workbook

from stock_assist.core import review, workbook_bytes
from stock_assist.filing import MAP_HEADERS, mapping_rows
from tests.test_core import file_for
from tests.test_groww import payload, stock_book, stock_row, mf_book, mf_row


def mapped(data):
    return [dict(zip(MAP_HEADERS, row)) for row in mapping_rows(data)]


@pytest.mark.parametrize(("sale", "period", "suffix"), [
    ("2024-07-22", "BE", "EquityMFonSTTDtls_BE"),
    ("2024-07-23", "AE", "EquityMFonSTTDtls"),
])
def test_2024_transfer_boundary_exact(sale, period, suffix):
    data = review([file_for(buy_date="2024-04-01", sell_date=sale)], "2024-25")
    assert data["filing"]["groups"][0]["period"] == period
    assert all(row["official_container"].endswith(suffix) for row in mapped(data))
    assert all(row["transfer_period"] == period for row in mapped(data))


@pytest.mark.parametrize("fy,sale,period", [
    ("2024-25", "2024-07-22", "BE"), ("2024-25", "2024-07-23", "AE"),
    ("2025-26", "2025-07-23", "FULL_FY"),
])
def test_112a_uses_only_year_applicable_transfer_flag(fy, sale, period):
    data = review([file_for(buy_date="2020-04-01", sell_date=sale)], fy)
    rows = mapped(data)
    flag = [r for r in rows if r["official_field"] == "ShareTransferredOnOrBefore"]
    assert len(flag) == (1 if fy == "2024-25" else 0)
    if flag:
        assert flag[0]["value"] == period
    assert next(r for r in rows if r["official_field"] == "ShareOnOrBefore")["value"] == "AE"
    assert all(r["schedule"].startswith("Schedule 112A") for r in rows)


def test_unknown_native_fields_are_blank_and_scheme_code_is_not_isin():
    data = review([payload(mf_book({"Equity Category": [mf_row()]}))], "2025-26")
    rows = mapped(data)
    assert data["included"] == 0
    assert all(r["status"] != "READY_WORKING" for r in rows)
    assert next(r for r in rows if r["field_id"] == "112a.isin")["value"] is None
    assert all(r["scheme_code"] == "SYNTH123" for r in rows)
    assert next(r for r in rows if r["field_id"] == "cg.transfer_expenses")["value"] is None
    assert next(r for r in rows if r["field_id"] == "cg.balance_before_adjustments")["value"] is None
    book = load_workbook(io.BytesIO(workbook_bytes(data)))
    assert book["Filing groups"]["H2"].value is None
    assert book["Filing review"]["B4"].value == "INCOMPLETE"
    assert book["Filing review"]["B3"].value == "2026-27"


def test_loss_and_profit_are_separate_no_threshold_or_double_counting():
    profit = file_for(cost_inr="0.10", proceeds_inr="0.30", transfer_expenses_inr="0.01")
    loss = file_for(lot_id="loss", cost_inr="1.00", proceeds_inr="0.20", transfer_expenses_inr="0")
    loss["name"] = "loss.csv"
    native = payload(stock_book({"Short Term trades": [stock_row()]}))
    data = review([profit, loss, native], "2025-26")
    groups = data["filing"]["groups"]
    assert len(groups) == 2
    canonical = next(g for g in groups if g["basis"] == "ATTESTED_CANONICAL")
    assert (canonical["positive_gains"], canonical["losses"], canonical["net"]) == ("0.19", "-0.80", "-0.61")
    assert data["totals"]["STCG"] == "-0.61"
    assert next(g for g in groups if g["basis"] == "BROKER_REPORTED")["status"] == "PROVISIONAL"
    assert all(r["value"] is None for r in mapped(data) if r["field_id"] in (
        "cg.section94_adjustment", "cg.capital_gain_after_adjustments",
    ))
    assert not any(r["schedule"] in ("Schedule SI", "CYLA", "BFLA", "CFL") for r in mapped(data))


def test_zero_differs_from_unknown_blocked_and_empty_are_not_nil_returns():
    data = review([file_for(transfer_expenses_inr="0")], "2025-26")
    expenses = next(r for r in mapped(data) if r["field_id"] == "cg.transfer_expenses")
    assert expenses["value"] == "0" and expenses["status"] == "READY_WORKING"
    blocked = review([file_for(transfer_expenses_inr="")], "2025-26")
    assert blocked["filing"]["groups"] == []
    assert mapped(blocked)[0]["status"] == "BLOCKED"
    assert mapped(blocked)[0]["official_field"] == ""
    assert mapped(blocked)[0]["value"] is None
    empty = review([payload(stock_book())], "2025-26")
    assert empty["filing"]["groups"] == []
    assert mapped(empty) == []
    book = load_workbook(io.BytesIO(workbook_bytes(empty)))
    assert book["Filing map"].max_row == 1
    assert book["Filing review"]["B4"].value == "EMPTY"


def test_exact_native_products_and_stable_snapshot_provenance_roundtrip():
    data = review([payload(mf_book({"Equity Category": [mf_row()]}), "synthetic.xlsx")], "2025-26")
    with localcontext() as ctx:
        ctx.prec = 6
        content = workbook_bytes(data)
    book = load_workbook(io.BytesIO(content))
    values = list(book["Filing map"].values)
    rows = [dict(zip(values[0], row)) for row in values[1:]]
    proceeds = next(r for r in rows if r["field_id"] == "cg.consideration")
    assert proceeds["value"] == data["records"][0]["raw"]["proceeds_inr"]
    assert proceeds["source_file"] == "synthetic.xlsx"
    assert proceeds["source_sheet"] == data["records"][0]["source_sheet"]
    assert proceeds["source_row"] == data["records"][0]["source_line"]
    assert proceeds["format_version"] == "1.0.0"
    assert proceeds["record_id"] == "lot-000001"
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    assert sum(Decimal(g["net"]) for g in data["filing"]["groups"]) == Decimal(data["broker_totals"]["LTCG"])


def test_no_ltcg_exemption_deducted_from_lot_or_group():
    data = review([file_for(buy_date="2020-01-01", cost_inr="0", proceeds_inr="200000", transfer_expenses_inr="0")], "2025-26")
    assert data["filing"]["groups"][0]["net"] == "200000.00"
    assert next(r for r in mapped(data) if r["official_field"] == "Balance")["value"] == "200000.00"


def test_2026_stcg_never_uses_before_flag_and_retains_official_spelling():
    data = review([file_for()], "2025-26")
    assert all(r["official_container"].endswith("EquityMFonSTTDtls") for r in mapped(data))
    assert next(r for r in mapped(data) if r["field_id"] == "cg.acquisition_cost")["official_field"] == "DeductSec48.AquisitCost"
