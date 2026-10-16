"""Synthetic workbooks mirror verified headers; no private source files used."""

import base64
import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from stock_assist.core import ImportProblem, review, workbook_bytes
from stock_assist.groww import MF_HEADERS, STOCK_HEADERS


def payload(book, name="capital-gains.xlsx"):
    stream = io.BytesIO()
    book.save(stream)
    return {"name": name, "base64": base64.b64encode(stream.getvalue()).decode()}


def stock_book(rows=None, offset=7, title="Capital Gains Statement for stocks"):
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic stock CG"
    sheet.append([title])
    for _ in range(offset):
        sheet.append([])
    for section in ("Intraday trades", "Short Term trades", "Long Term trades", "Other trades"):
        sheet.append([section])
        sheet.append(STOCK_HEADERS)
        for row in (rows or {}).get(section, []):
            sheet.append(row)
        sheet.append([])
    return book


def stock_row(**overrides):
    values = dict(zip(STOCK_HEADERS, [
        "Synthetic listed equity", "INE000A01001", 2, "10-05-2025", 100.05,
        200.10, "10-12-2025", 150.15, 300.30, 100.20, "",
    ]))
    values.update(overrides)
    return [values[key] for key in STOCK_HEADERS]


def mf_book(rows=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic fund CG"
    sheet.append(["Mutual funds capital gains"])
    for section in ("Equity Category", "Debt (Specified - Other than Equity) Category",
                    "Debt (Unspecified - Other than Equity) Category"):
        sheet.append([section])
        sheet.append(MF_HEADERS)
        for row in (rows or {}).get(section, []):
            sheet.append(row)
        sheet.append([])
    return book


def mf_row(**overrides):
    values = dict(zip(MF_HEADERS, [
        "Synthetic equity fund", "SYNTH123", "Equity", "synthetic-folio",
        "purchase-1", "2023-06-01", 1.234567, 100.123456, "redeem-1", "2025-07-01",
        10, 200.123456, 0, 123.46,
    ]))
    values.update(overrides)
    return [values[key] for key in MF_HEADERS]


def test_stock_native_review_export_and_traceability():
    file = payload(stock_book({"Short Term trades": [stock_row()]}))
    data = review([file], "2025-26")
    assert data["broker_totals"] == {"STCG": "100.20", "LTCG": "0.00"}
    assert data["totals"] == {"STCG": "0.00", "LTCG": "0.00"}
    assert data["broker_review"] == 1
    assert data["blocked"] == 1 and data["included"] == 0
    record = data["records"][0]
    assert record["status"] == "REVIEW_REQUIRED"
    assert record["source_sheet"] == "Synthetic stock CG"
    assert record["source_line"] == 14
    assert record["gross_gain"] == "100.2"
    assert record["raw"]["transfer_expenses_inr"] == ""
    assert record["raw"]["stt_paid"] == ""
    book = load_workbook(io.BytesIO(workbook_bytes(data)))
    assert book["Disposed lots"]["V2"].value == "Synthetic stock CG"
    assert book["Disposed lots"]["Y2"].value == "100.2"
    assert book["Import manifest"]["D2"].value == 1


def test_mf_unit_products_and_scheme_identity_without_isin():
    data = review([payload(mf_book({"Equity Category": [mf_row()]}))], "2025-26")
    r = data["records"][0]
    assert r["raw"]["isin"] == ""
    assert r["scheme_code"] == "SYNTH123"
    assert r["raw"]["cost_inr"] == "123.609114703552"
    assert r["gross_gain"] == "123.456700000000"
    assert data["broker_totals"]["LTCG"] == "123.46"
    assert data["broker_review"] == 1
    assert "Folio Number" not in r["source_values"]
    assert r["source_values"]["Grandfathered Nav"] == "10"


def test_intraday_unknown_debt_and_out_of_year_remain_visible_excluded():
    stock = stock_book({
        "Intraday trades": [stock_row(**{"Buy date": "10-12-2025"})],
        "Short Term trades": [stock_row(**{"Sell date": "10-12-2026"})],
        "Other trades": [stock_row()],
    })
    funds = mf_book({
        "Debt (Specified - Other than Equity) Category": [mf_row(Category="Debt")],
        "Equity Category": [mf_row(Category="Unknown")],
    })
    data = review([payload(stock, "stock.xlsx"), payload(funds, "mf.xlsx")], "2025-26")
    assert len(data["records"]) == 5
    assert data["broker_review"] == 0
    assert all(r["status"] == "BLOCKED" for r in data["records"])
    assert all(r["gain"] == "" for r in data["records"])


@pytest.mark.parametrize("overrides", [
    {"Realised P&L": 99.99}, {"Remark": "Review adjusted acquisition"},
    {"Buy date": "2017-01-01"}, {"Buy value": None}, {"Quantity": 0},
    {"Buy date": "31-02-2025"}, {"Realised P&L": None}, {"ISIN": ""},
    {"Realised P&L": 100.201},
])
def test_stock_missing_ambiguous_or_unreconciled_blocks(overrides):
    data = review([payload(stock_book({"Short Term trades": [stock_row(**overrides)]}))], "2025-26")
    assert data["broker_review"] == 0
    assert data["records"][0]["status"] == "BLOCKED"
    assert data["records"][0]["issues"]


def test_wrong_reported_term_blocks():
    data = review([payload(stock_book({"Long Term trades": [stock_row()]}))], "2025-26")
    assert data["broker_review"] == 0
    assert any("classification differs" in issue for issue in data["records"][0]["issues"])


def test_empty_reports_are_explicit_not_failed_or_fabricated():
    data = review([payload(stock_book(), "stock.xlsx"), payload(mf_book(), "mf.xlsx")], "2025-26")
    assert data["records"] == []
    assert all(m["outcome"].startswith("EMPTY") for m in data["imports"])
    book = load_workbook(io.BytesIO(workbook_bytes(data)))
    assert book["Summary"]["B4"].value.startswith("EMPTY")
    assert book["Import manifest"].max_row == 3


def test_duplicates_across_reports_block_all_copies_but_not_identical_lots_within_one():
    rows = {"Short Term trades": [stock_row(), stock_row()]}
    first = payload(stock_book(rows), "first.xlsx")
    second = payload(stock_book(rows, offset=20), "second.xlsx")
    data = review([first], "2025-26")
    assert data["broker_review"] == 2
    combined = review([first, second], "2025-26")
    assert len(combined["records"]) == 4
    assert combined["broker_review"] == 0
    assert all(any("Duplicate" in i for i in r["issues"]) for r in combined["records"])


def test_duplicate_identity_ignores_numeric_display_and_date_cell_formatting():
    first = payload(stock_book({"Short Term trades": [stock_row()]}), "first.xlsx")
    second = payload(stock_book({"Short Term trades": [stock_row(**{
        "Stock name": "Changed display name", "Quantity": "2.0000", "Buy value": "200.100",
        "Sell value": "300.300", "Buy date": datetime(2025, 5, 10),
    })]}), "second.xlsx")
    data = review([first, second], "2025-26")
    assert data["broker_review"] == 0
    assert all(any("Duplicate" in issue for issue in r["issues"]) for r in data["records"])


def test_unit_products_keep_more_than_28_digits_of_precision():
    fund = mf_row(**{
        "Matched Quantity": "999999999999.123456",
        "Purchase Price": "1234567890123456.123456789012",
        "Redeem Price": "1234567890123456.123456789013",
        "Long Term-Capital Gain": "1.00",
    })
    data = review([payload(mf_book({"Equity Category": [fund]}))], "2025-26")
    assert data["records"][0]["gross_gain"] == "0.999999999999123456"
    assert data["broker_totals"]["LTCG"] == "1.00"


@pytest.mark.parametrize("overrides", [
    {"Category": "Debt"}, {"Scheme Code": ""}, {"Purchase Transaction Id": ""},
    {"Long Term-Capital Gain": 120}, {"Short Term-Capital Gain": 1},
    {"Purchase Date": "2018-01-31"}, {"Purchase Price": None},
])
def test_mf_ambiguities_are_excluded(overrides):
    data = review([payload(mf_book({"Equity Category": [mf_row(**overrides)]}))], "2025-26")
    assert data["broker_review"] == 0
    assert data["records"][0]["status"] == "BLOCKED"


def test_loss_zero_and_excel_date_values():
    row = stock_row(**{"Buy date": datetime(2025, 5, 10), "Sell date": datetime(2025, 12, 10),
                       "Sell value": 100.10, "Realised P&L": -100})
    zero = stock_row(**{"ISIN": "INE000B01002", "Sell value": 200.10, "Realised P&L": 0})
    data = review([payload(stock_book({"Short Term trades": [row, zero]}))], "2025-26")
    assert data["broker_totals"]["STCG"] == "-100.00"
    assert data["broker_review"] == 2


def test_reject_pnl_by_name_title_or_multiple_sheets():
    for book, name in [
        (stock_book(), "Stocks_PnL_Report.xlsx"),
        (stock_book(title="P&L Statement for stocks"), "renamed.xlsx"),
    ]:
        with pytest.raises(ImportProblem, match="P&L reports"):
            review([payload(book, name)], "2025-26")
    book = stock_book()
    book.create_sheet("Aggregated holdings")
    with pytest.raises(ImportProblem, match="single-sheet"):
        review([payload(book)], "2025-26")


def test_unrecognized_truncated_and_formula_workbooks_are_rejected():
    for book in [Workbook(), stock_book({"Short Term trades": [stock_row(**{"Buy value": "=1+2"})]})]:
        with pytest.raises(ImportProblem):
            review([payload(book)], "2025-26")
    with pytest.raises(ImportProblem, match="Cannot read"):
        review([{"name": "bad.xlsx", "base64": base64.b64encode(b"not xlsx").decode()}], "2025-26")
    with pytest.raises(ImportProblem, match="Invalid XLSX"):
        review([{"name": "bad.xlsx", "base64": "not base64"}], "2025-26")


def test_changed_header_and_unknown_section_do_not_silently_match():
    book = stock_book({"Short Term trades": [stock_row()]})
    for row in book.active:
        if row[0].value == "Stock name":
            row[1].value = "Security ID"
            break
    with pytest.raises(ImportProblem, match="changed Groww header"):
        review([payload(book)], "2025-26")
    book = stock_book({"Short Term trades": [stock_row()]})
    for row in book.active:
        if row[0].value == "Short Term trades":
            row[0].value = "Unverified asset class"
            break
    data = review([payload(book)], "2025-26")
    assert data["broker_review"] == 0
    assert len(data["records"]) == 1


def test_out_of_year_and_debt_have_no_computed_equity_term():
    stock = payload(stock_book({"Short Term trades": [stock_row(**{"Sell date": "10-12-2026"})]}))
    fund = payload(mf_book({"Debt (Specified - Other than Equity) Category": [mf_row(Category="Debt")]}), "mf.xlsx")
    data = review([stock, fund], "2025-26")
    assert all(r["term"] == "" for r in data["records"])
    assert data["records"][0]["broker_term"] == "STCG"
