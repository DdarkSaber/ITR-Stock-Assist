import csv
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from stock_assist.core import HEADERS, ImportProblem, parse_csv, review, workbook_bytes

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = (ROOT / "samples" / "canonical-demo.csv").read_text(encoding="utf-8")


def file_for(**changes):
    raw = next(csv.DictReader(io.StringIO(SAMPLE)))
    raw.update(changes)
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, HEADERS)
    writer.writeheader()
    writer.writerow(raw)
    return {"name": "test.csv", "text": out.getvalue()}


def test_sample_import_to_xlsx():
    data = review([{"name": "demo.csv", "text": SAMPLE}], "2025-26")
    assert data["totals"] == {"STCG": "1965.20", "LTCG": "7500.00"}
    assert (data["included"], data["blocked"]) == (3, 1)
    assert data["ay"] == "2026-27"
    book = load_workbook(io.BytesIO(workbook_bytes(data)))
    assert book["Summary"]["B1"].value == "Capital-gains working paper v0.3.1"
    assert book.sheetnames == [
        "Filing review", "Filing groups", "Filing map", "References",
        "Summary", "Disposed lots", "Review issues", "Source rows", "Import manifest",
    ]
    assert book["Source rows"].max_row == 5
    assert book["Summary"]["B7"].value == "1965.20"
    assert book["Summary"]["B7"].data_type == "s"
    assert book["Disposed lots"]["B5"].value == 5
    assert book["Review issues"].max_row == 5
    assert book["Summary"]["B4"].value.startswith("INCOMPLETE")


@pytest.mark.parametrize(("buy", "sell", "term"), [
    ("2024-05-10", "2025-05-10", "LTCG"),
    ("2024-05-10", "2025-05-09", "STCG"),
    ("2024-05-10", "2025-05-11", "LTCG"),
    ("2024-02-29", "2025-03-01", "LTCG"),
])
def test_calendar_holding_period(buy, sell, term):
    fy = "2024-25" if sell < "2025-04-01" else "2025-26"
    data = review([file_for(buy_date=buy, sell_date=sell)], fy)
    assert data["records"][0]["term"] == term


@pytest.mark.parametrize("changes", [
    {"cost_inr": ""}, {"cost_inr": "NaN"}, {"cost_inr": "Infinity"},
    {"cost_inr": "1.001"}, {"cost_inr": "-1"}, {"cost_inr": "1,000"},
    {"cost_inr": "1e3"}, {"cost_inr": "1000000000000"},
    {"sell_date": "2025-02-29"}, {"sell_date": "2025-5-01"},
    {"buy_date": "2025-12-10"}, {"buy_date": "2026-01-01"},
    {"asset_type": "debt_mf"}, {"stt_paid": "no"},
    {"corporate_action": "split"}, {"cost_basis_confirmed": "no"},
    {"buy_date": "2018-01-31"}, {"quantity": "0"}, {"quantity": "1.0000001"},
    {"isin": "INVALID"}, {"broker": ""}, {"lot_id": ""},
    {"buy_date": "2024-02-29", "sell_date": "2025-02-28"},
])
def test_unsupported_or_missing_data_never_calculated(changes):
    data = review([file_for(**changes)], "2025-26")
    assert data["included"] == 0
    assert data["records"][0]["issues"]
    assert data["records"][0]["gain"] == ""
    assert data["totals"] == {"STCG": "0.00", "LTCG": "0.00"}


@pytest.mark.parametrize(("sell", "included"), [
    ("2025-03-31", 0), ("2025-04-01", 1), ("2026-03-31", 1), ("2026-04-01", 0),
])
def test_financial_year_boundaries(sell, included):
    data = review([file_for(buy_date="2024-04-01", sell_date=sell)], "2025-26")
    assert data["included"] == included


def test_duplicates_block_all_copies_across_files():
    first = file_for()
    second = {**first, "name": "second.csv"}
    data = review([first, second], "2025-26")
    assert data["blocked"] == 2
    assert all("Duplicate" in " ".join(r["issues"]) for r in data["records"])
    with pytest.raises(ImportProblem, match="names must be unique"):
        review([first, first], "2025-26")


def test_exact_decimal_loss_and_zero_cost():
    data = review([file_for(cost_inr="0.10", proceeds_inr="0.30", transfer_expenses_inr="0.01")], "2025-26")
    assert data["totals"]["STCG"] == "0.19"
    data = review([file_for(cost_inr="0", proceeds_inr="0", transfer_expenses_inr="0.10")], "2025-26")
    assert data["totals"]["STCG"] == "-0.10"
    assert data["records"][0]["status"] == "INCLUDED"
    assert any("likely rights entitlement or corporate action - confirm" in issue
               for issue in data["records"][0]["issues"])


@pytest.mark.parametrize("text", ["", "a,b\n1,2\n", SAMPLE.replace("lot_id", "broker", 1),
                                      SAMPLE + "one,two\n", SAMPLE + '"unclosed'])
def test_malformed_csv_is_atomic(text):
    with pytest.raises(ImportProblem):
        parse_csv(text, "broken.csv")


def test_utf8_bom_quoted_fields_and_source_line():
    f = file_for(asset_name='Synthetic, "quoted"\nname')
    data = review([{**f, "text": "\ufeff" + f["text"]}], "2025-26")
    assert data["records"][0]["source_line"] == 3
    assert data["records"][0]["raw"]["asset_name"] == 'Synthetic, "quoted"\nname'


def test_spreadsheet_formula_injection_is_literal():
    data = review([file_for(asset_name='=HYPERLINK("https://example.invalid","click")')], "2025-26")
    book = load_workbook(io.BytesIO(workbook_bytes(data)), data_only=False)
    for sheet in book:
        for row in sheet:
            assert all(cell.data_type != "f" for cell in row)
    assert book["Source rows"]["E2"].value.startswith("=HYPERLINK")


def test_zero_cost_and_sale_are_flagged_as_possible_rights_or_corporate_action():
    data = review([file_for(cost_inr="0", proceeds_inr="0", transfer_expenses_inr="0")], "2025-26")
    record = data["records"][0]
    assert record["status"] == "BLOCKED"
    assert "likely rights entitlement or corporate action - confirm" in " ".join(record["issues"])
    book = load_workbook(io.BytesIO(workbook_bytes(data)))
    assert any("likely rights entitlement or corporate action - confirm" in row[3].value
               for row in book["Review issues"].iter_rows(min_row=2))


def test_file_and_year_limits():
    with pytest.raises(ImportProblem, match="only canonical"):
        parse_csv(SAMPLE, "groww.xlsx")
    with pytest.raises(ImportProblem, match="2 MB"):
        parse_csv("x" * 2_000_001, "large.csv")
    with pytest.raises(ImportProblem, match="Choose FY"):
        review([file_for()], "2026-27")
    with pytest.raises(ImportProblem, match="between 1 and 10"):
        review([], "2025-26")
    with pytest.raises(ImportProblem, match="control characters"):
        parse_csv(file_for(asset_name="bad\x00name")["text"], "bad.csv")


def test_too_many_rows():
    row = SAMPLE.splitlines()[1]
    text = SAMPLE.splitlines()[0] + "\n" + (row + "\n") * 10001
    with pytest.raises(ImportProblem, match="maximum 10000"):
        parse_csv(text, "too-many.csv")
