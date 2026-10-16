"""Adapters for verified Groww capital-gains layouts, not P&L statements."""

import base64
import binascii
import io
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from stock_assist.core import HEADERS, ImportProblem, MAX_BYTES, MAX_ROWS

STOCK_HEADERS = (
    "Stock name", "ISIN", "Quantity", "Buy date", "Buy price", "Buy value",
    "Sell date", "Sell price", "Sell value", "Realised P&L", "Remark",
)
MF_HEADERS = (
    "Scheme Name", "Scheme Code", "Category", "Folio Number",
    "Purchase Transaction Id", "Purchase Date", "Matched Quantity", "Purchase Price",
    "Redeem Transaction Id", "Redeem Date", "Grandfathered Nav", "Redeem Price",
    "Short Term-Capital Gain", "Long Term-Capital Gain",
)
SECTIONS = {
    "intraday trades": "intraday", "short term trades": "STCG",
    "long term trades": "LTCG", "others": "other", "others trades": "other",
    "other trades": "other", "equity category": "equity_mf",
    "debt (specified - other than equity) category": "debt",
    "debt (unspecified - other than equity) category": "debt",
}
PENDING = (
    "Broker report only: verify STT, acquisition/corporate-action history and "
    "allowable costs/fees. Transfer expenses are not supplied per lot; no zero "
    "expense or tax-eligibility confirmation has been invented. Use a verified "
    "canonical CSV for working-paper inclusion."
)


def _text(value):
    return "" if value is None else str(value).strip()


def _decimal(value, field, signed=False):
    text = _text(value)
    if not re.fullmatch(r"-?\d{1,16}(?:\.\d{1,12})?", text):
        raise ValueError(f"{field}: missing or unsupported numeric value.")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{field}: invalid decimal.") from exc
    if not signed and number < 0:
        raise ValueError(f"{field}: negative value needs manual review.")
    return number


def _iso(value, field):
    if isinstance(value, datetime):
        if value.time() != datetime.min.time():
            raise ValueError(f"{field}: non-midnight timestamp requires review.")
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    for pattern, fmt in ((r"\d{2}-\d{2}-\d{4}", "%d-%m-%Y"),
                         (r"\d{4}-\d{2}-\d{2}", "%Y-%m-%d")):
        if re.fullmatch(pattern, text):
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError as exc:
                raise ValueError(f"{field}: invalid calendar date.") from exc
    raise ValueError(f"{field}: expected DD-MM-YYYY or YYYY-MM-DD.")


def _fixed(number):
    return format(number, "f")


def _identity_number(value):
    return format(Decimal(value).normalize(), "f") if value else ""


def parse_groww(encoded, filename):
    """Return records plus an explicit import manifest, even for empty reports."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from xml.etree.ElementTree import ParseError

    if "pnl" in filename.casefold() or "p&l" in filename.casefold():
        raise ImportProblem("Groww P&L reports are not supported. Use the Stock Capital "
                            "Gains report instead; importing both would double-count disposals.")
    if len(encoded) > 4 * ((MAX_BYTES + 2) // 3):
        raise ImportProblem("XLSX exceeds the 2 MB file limit.")
    try:
        payload = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ImportProblem("Invalid XLSX data. Select the original Groww capital-gains workbook.") from exc
    if len(payload) > MAX_BYTES:
        raise ImportProblem("XLSX exceeds the 2 MB file limit.")
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            if len(archive.infolist()) > 500 or sum(i.file_size for i in archive.infolist()) > 30_000_000:
                raise ImportProblem("XLSX archive exceeds safe expansion limits.")
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, InvalidFileException, ParseError, KeyError, OSError, ValueError) as exc:
        raise ImportProblem("Cannot read this XLSX. Select an unmodified Groww capital-gains report.") from exc
    try:
        if len(book.worksheets) != 1:
            raise ImportProblem("Expected a single-sheet Groww capital-gains report. "
                                "Use the Stock or Mutual Fund Capital Gains report; P&L "
                                "and aggregated holdings workbooks would risk double counting.")
        sheet = book.worksheets[0]
        if not sheet.max_row or not sheet.max_column or sheet.max_row > 15000 or sheet.max_column > 30:
            raise ImportProblem("XLSX dimensions exceed the supported Groww report limits.")
        rows = []
        for cells in sheet.iter_rows():
            if any(c.data_type == "f" for c in cells):
                raise ImportProblem("Formula cells are unsupported. Use the original downloaded "
                                    "Groww capital-gains report, not an edited calculation workbook.")
            values = [c.value for c in cells]
            if any(len(_text(v)) > 2000 for v in values):
                raise ImportProblem("XLSX contains a field longer than 2,000 characters.")
            rows.append(values)
        labels = [_text(v) for row in rows[:50] for v in row]
        if any("p&l statement" in v.casefold() or "unrealised trades" in v.casefold() for v in labels):
            raise ImportProblem("Groww P&L reports are not supported. Use the Stock Capital "
                                "Gains report instead to avoid double counting.")
        stock = any(tuple(_text(v) for v in row[:11]) == STOCK_HEADERS for row in rows)
        mf = any(tuple(_text(v) for v in row[:14]) == MF_HEADERS for row in rows)
        if stock == mf:
            raise ImportProblem("Unrecognized or mixed XLSX schema. Supported: original Groww "
                                "Stock Capital Gains and Mutual Fund Capital Gains reports.")
        kind = "groww_stock_cg" if stock else "groww_mf_cg"
        expected = STOCK_HEADERS if stock else MF_HEADERS
        section = None
        active = False
        headers_seen = 0
        records = []
        for line, values in enumerate(rows, 1):
            values += [None] * max(0, len(expected) - len(values))
            first = _text(values[0])
            if not first and not any(v is not None for v in values):
                continue
            if first.casefold() in SECTIONS:
                section = SECTIONS[first.casefold()]
                active = False
                continue
            if tuple(_text(v) for v in values[:len(expected)]) == expected:
                headers_seen += 1
                active = True
                continue
            if first in ("Stock name", "Scheme Name"):
                raise ImportProblem(f"Sheet row {line}: changed Groww header layout; "
                                    "this report version needs a verified adapter update.")
            if not active:
                continue
            # Data rows always have an asset name and populated lot fields. Totals and
            # known empty-table messages are scaffolding, not disposals.
            if not first:
                if any(v not in (None, "") for v in values[1:]):
                    raise ImportProblem(f"Sheet row {line}: unnamed row inside a lot table; "
                                        "cannot safely identify a disposal.")
                continue
            if first.casefold() in ("total", "grand total", "no transactions", "no trades",
                                    "no transactions found", "no trades found"):
                continue
            # A merged section label ends a table, including sections we do not recognize.
            if all(v in (None, "") for v in values[1:]):
                section = "other"
                active = False
                continue
            original = dict(zip(expected, values))
            record = _stock(original, section) if stock else _mf(original, section)
            record.update(source_file=filename, source_line=line, source_sheet=sheet.title,
                          source_format=kind, source_values={
                              k: _text(v) for k, v in original.items() if k != "Folio Number"
                          })
            records.append(record)
            if len(records) > MAX_ROWS:
                raise ImportProblem("Groww report exceeds 10,000 disposed lots.")
        # Verified exports always contain separate tables, including empty sections.
        if headers_seen < 3:
            raise ImportProblem("Incomplete Groww report: expected all capital-gains section tables.")
        occurrences = Counter()
        for record in records:
            identity = record["identity"]
            occurrences[identity] += 1
            record["duplicate_key"] = (kind, identity, occurrences[identity])
        return records, {
            "source_file": filename, "source_sheet": sheet.title, "format": kind,
            "rows": len(records),
            "outcome": "EMPTY: no disposal rows reported" if not records else "Imported for broker reconciliation",
        }
    except ImportProblem:
        raise
    except (zipfile.BadZipFile, ParseError, KeyError, OSError, ValueError) as exc:
        raise ImportProblem("Cannot read report cells. Download a fresh original Groww "
                            "capital-gains XLSX and retry.") from exc
    finally:
        book.close()


def _base():
    return {"raw": dict.fromkeys(HEADERS, ""), "adapter_issues": [],
            "broker_gain": "", "broker_term": "", "gross_gain": "",
            "reconciliation": "", "identity": (), "native": True}


def _numbers(record, values, quantity, buy, sell, gain):
    issues = record["adapter_issues"]
    parsed = {}
    for field in (quantity, buy, sell, *gain):
        try:
            parsed[field] = _decimal(values[field], field, signed=field in gain)
            if field in gain and parsed[field] != parsed[field].quantize(Decimal(".01")):
                issues.append(f"{field}: broker gain has sub-paise precision; excluded "
                              "rather than silently rounded.")
        except ValueError as exc:
            issues.append(str(exc))
    if quantity in parsed:
        if parsed[quantity] <= 0:
            issues.append("Disposed quantity must be positive.")
        record["raw"]["quantity"] = _fixed(parsed[quantity])
    return parsed


def _dates(record, values, buy, sell):
    for source, target in ((buy, "buy_date"), (sell, "sell_date")):
        try:
            record["raw"][target] = _iso(values[source], source)
        except ValueError as exc:
            record["adapter_issues"].append(str(exc))


def _reconcile(record, cost, proceeds, broker_gain):
    record["raw"]["cost_inr"] = _fixed(cost)
    record["raw"]["proceeds_inr"] = _fixed(proceeds)
    gross = proceeds - cost
    record["gross_gain"] = _fixed(gross)
    if broker_gain is not None:
        record["broker_gain"] = _fixed(broker_gain)
        difference = gross - broker_gain
        record["reconciliation"] = (
            "MATCH within INR 0.01" if abs(difference) <= Decimal(".01")
            else f"MISMATCH: gross minus reported = {_fixed(difference)} INR"
        )
        if abs(difference) > Decimal(".01"):
            record["adapter_issues"].append("Broker gain does not reconcile to gross proceeds "
                                             "minus cost within INR 0.01.")


def _stock(values, section):
    r = _base()
    r["raw"].update(broker="Groww", asset_name=_text(values["Stock name"]),
                    isin=_text(values["ISIN"]),
                    asset_type="listed_equity" if section in ("STCG", "LTCG") else "unknown")
    r["broker_term"] = section if section in ("STCG", "LTCG") else ""
    if section not in ("STCG", "LTCG"):
        r["adapter_issues"].append(f"Unsupported stock report section: {section or 'unknown'}; "
                                   "intraday/other trades are excluded.")
    if _text(values["Remark"]) not in ("", "-", "--"):
        r["adapter_issues"].append("Broker remark present; inspect original source for cost/date "
                                   "or corporate-action qualifications.")
    _dates(r, values, "Buy date", "Sell date")
    numbers = _numbers(r, values, "Quantity", "Buy value", "Sell value", ("Realised P&L",))
    if "Realised P&L" in numbers:
        r["broker_gain"] = _fixed(numbers["Realised P&L"])
    if "Buy value" in numbers and "Sell value" in numbers:
        _reconcile(r, numbers["Buy value"], numbers["Sell value"], numbers.get("Realised P&L"))
    r["identity"] = (
        r["raw"]["isin"], r["raw"]["buy_date"], r["raw"]["sell_date"],
        *(_identity_number(r["raw"][key]) for key in ("quantity", "cost_inr", "proceeds_inr")),
    )
    r["raw"]["lot_id"] = "stock:" + "|".join(r["identity"])
    return r


def _mf(values, section):
    r = _base()
    equity = section == "equity_mf" and _text(values["Category"]) == "Equity"
    scheme = _text(values["Scheme Code"])
    r["scheme_code"] = scheme
    r["raw"].update(broker="Groww", asset_name=_text(values["Scheme Name"]),
                    asset_type="equity_mf" if equity else "unknown")
    if not equity:
        r["adapter_issues"].append("Debt or unknown fund category: excluded from equity calculations.")
    if not scheme:
        r["adapter_issues"].append("Missing scheme code; fund identity is not established.")
    if not _text(values["Purchase Transaction Id"]) or not _text(values["Redeem Transaction Id"]):
        r["adapter_issues"].append("Missing purchase/redemption transaction ID.")
    _dates(r, values, "Purchase Date", "Redeem Date")
    numbers = _numbers(r, values, "Matched Quantity", "Purchase Price", "Redeem Price",
                       ("Short Term-Capital Gain", "Long Term-Capital Gain"))
    short = numbers.get("Short Term-Capital Gain")
    long = numbers.get("Long Term-Capital Gain")
    if short is not None and long is not None:
        r["broker_gain"] = _fixed(short + long)
        r["broker_term"] = "STCG" if short != 0 and long == 0 else (
            "LTCG" if long != 0 and short == 0 else "")
        if short != 0 and long != 0:
            r["adapter_issues"].append("Both STCG and LTCG populated on one matched lot; review required.")
    if all(k in numbers for k in ("Matched Quantity", "Purchase Price", "Redeem Price")):
        quantity = numbers["Matched Quantity"]
        # Preserve exact unit products; do not round each side independently.
        _reconcile(r, quantity * numbers["Purchase Price"], quantity * numbers["Redeem Price"],
                   short + long if short is not None and long is not None else None)
    # Groww supplies this reference NAV even for recent purchases. Preserve it as
    # source evidence, but never substitute it for actual cost; old dates are blocked by core.
    r["identity"] = (
        scheme, _text(values["Purchase Transaction Id"]), _text(values["Redeem Transaction Id"]),
        _identity_number(r["raw"]["quantity"]), r["raw"]["buy_date"], r["raw"]["sell_date"],
    )
    r["raw"]["lot_id"] = "mf:" + "|".join(r["identity"])
    return r
